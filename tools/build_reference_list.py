"""Development-time builder of a private reference list from a reference-peaks workbook.

Reads one sheet laid out as: a header row starting "Material", "Description", then
one column per band position (the header holds the position), and a "References"
column whose cell links to the source. Each material row marks its bands by filling
those columns. Writes a reference list (schema raman-studio/reference-list@1) that
the app can load for a session ("Load a reference list..." on the Peaks tab).

The list is private: the output must lie inside the gitignored test-data/private/
folder, and nothing is published. Names, roles and positions are copied verbatim;
a row without band positions or without a source link is left out and reported.
Never part of the deployed app; never run in CI. Needs openpyxl.

Usage:
    python tools/build_reference_list.py WORKBOOK.xlsx [--sheet Raman] [--tolerance 10]
        [--name "Battery materials (Raman)"] [--out test-data/private/references/NAME.json]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PRIVATE = (ROOT / "test-data" / "private").resolve()


def build(path: Path, sheet: str, tolerance: float, name: str):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"no sheet {sheet!r} in {path.name} (sheets: {', '.join(wb.sheetnames)})")
    ws = wb[sheet]
    header_row = next((r for r in ws.iter_rows() if str(r[0].value or "").strip().lower() == "material"), None)
    if header_row is None:
        raise SystemExit(f"sheet {sheet!r}: no header row starting with 'Material'")
    hr = header_row[0].row
    cols, ref_col, role_col = {}, None, None
    for c in header_row:
        v = c.value
        if isinstance(v, (int, float)):
            cols[c.column] = float(v)
        elif isinstance(v, str) and v.strip().lower().startswith("reference"):
            ref_col = c.column
        elif isinstance(v, str) and v.strip().lower() == "description":
            role_col = c.column
    materials, skipped, mismatched = [], [], []
    for row in ws.iter_rows(min_row=hr + 1):
        mat = row[0].value
        if not mat:
            continue
        bands = []
        for c in row:
            if c.column in cols and c.value is not None and str(c.value).strip() != "":
                v = float(c.value)
                if abs(v - cols[c.column]) > 1e-9:
                    mismatched.append(f"{mat}: {v} under the {cols[c.column]:g} column")
                bands.append(v)
        link = None
        if ref_col:
            cell = ws.cell(row=row[0].row, column=ref_col)
            if cell.hyperlink and (cell.hyperlink.target or "").lower().startswith(("http://", "https://")):
                link = {"url": cell.hyperlink.target, "text": str(cell.value or "").strip()}
        role = str(ws.cell(row=row[0].row, column=role_col).value or "").strip() if role_col else ""
        if not bands or not link:
            skipped.append(f"{mat}: {'no band positions' if not bands else 'no source link'}")
            continue
        materials.append({"name": str(mat).strip(), "role": role, "bands": sorted(bands), "source": link})
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    out = {
        "schema": "raman-studio/reference-list@1",
        "id": "private-" + path.stem.lower().replace(" ", "-")[:60],
        "name": name,
        "private": True,
        "toleranceCm": tolerance,
        "note": "Candidate positions from a private workbook; not reviewed for publication.",
        "source": {"file": path.name, "sheet": sheet, "sha256": digest, "builtWith": "tools/build_reference_list.py"},
        "materials": materials,
    }
    return out, skipped, mismatched


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--sheet", default="Raman")
    ap.add_argument("--tolerance", type=float, default=10.0, help="candidate tolerance in cm-1 (shown and adjustable in the app)")
    ap.add_argument("--name", default=None)
    ap.add_argument("--out", type=Path, default=None)
    a = ap.parse_args()
    if not (0 < a.tolerance <= 100):
        print("--tolerance must be between 0 and 100 cm-1")
        return 2
    out_path = (a.out or (PRIVATE / "references" / (a.workbook.stem + "-" + a.sheet.lower() + ".json"))).resolve()
    if PRIVATE not in out_path.parents:
        print(f"refusing to write {out_path}: reference lists built from private workbooks stay inside {PRIVATE}")
        return 2
    data, skipped, mismatched = build(a.workbook, a.sheet, a.tolerance, a.name or f"{a.workbook.stem} ({a.sheet})")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {out_path}: {len(data['materials'])} materials, tolerance ±{a.tolerance:g} cm-1")
    for s in skipped:
        print("  left out:", s)
    for m in mismatched:
        print("  value differs from its column header:", m)
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
